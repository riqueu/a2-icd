"""Módulo para exportar os dados para o excel"""

import pandas as pd
from openpyxl import load_workbook
from openpyxl.worksheet.table import Table, TableStyleInfo

#Sobrescreve as planilhas do excel com os dados brutos (em tabelas)
def update_tables(movies, reviews, box_offices):
    #importar o arquivo e atualizar os dados dos dataframes
    results = pd.read_excel('a2-icd\letterboxd_scraping\\results.xlsx', sheet_name=None)
    
    results['movies'] = movies
    results['reviews'] = reviews
    results['box offices'] = box_offices
    
    #atualizar o arquivo com os dataframes
    with pd.ExcelWriter('a2-icd\letterboxd_scraping\\results.xlsx', engine='openpyxl') as writer:
        for sheet_name, data in results.items():
            data.to_excel(writer, sheet_name=sheet_name, index=False)
            
    workbook = load_workbook('a2-icd\letterboxd_scraping\\results.xlsx')
    
    def create_table(sheet, title, ref):
        table = Table(displayName=title, ref=ref)
        style = TableStyleInfo(
            name="TableStyleMedium9",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=True,
        )
        table.tableStyleInfo = style
        sheet.add_table(table)
        
    #criar as tabelas com os dataframes
    for sheet_name in ['movies', 'reviews', 'box offices']:
        sheet = workbook[sheet_name]
        num_cols = results[sheet_name].shape[1]
        num_rows = results[sheet_name].shape[0]
        ref = f"A1:{chr(65 + num_cols - 1)}{num_rows + 1}"
        create_table(sheet, f"Table{sheet_name.title().replace(' ', '')}", ref)
    
    workbook.save('a2-icd\letterboxd_scraping\\results.xlsx')

def create_empty_table():
    with pd.ExcelWriter('a2-icd\letterboxd_scraping\\results.xlsx') as writer:
        empty_df = pd.DataFrame()
    
    # Add empty sheets
        empty_df.to_excel(writer, sheet_name='movies', index=False)
        empty_df.to_excel(writer, sheet_name='reviews', index=False)
        empty_df.to_excel(writer, sheet_name='box offices', index=False)